# -*- coding: utf-8 -*-
"""Informe de rivals: què fan els rivals quan juguen contra nosaltres.

Per a cada rival es compara el partit contra nosaltres amb la mitjana dels seus
altres partits (1 a 3). Després es fa la mitjana entre rivals, amb el mateix
pes per a tots, i se'n treu un PowerPoint de 10 diapositives.

Tot surt del play-by-play de la FCBQ (Cistella / Intent fallat de 1, 2 i 3), així
que no cal passar per l'Excel ni pateix el problema de les dates ("8/17" que
Excel convertia en agost del 2017).

No depèn de Streamlit: app.py li passa els DataFrames ja carregats.
"""
import io
import re

import pandas as pd
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import (XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION,
                             XL_TICK_LABEL_POSITION)
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt, Emu

# ══════════════════════════════════════════════════
# CÀLCULS
# ══════════════════════════════════════════════════

# Accions exactes del play-by-play (extraction.EVENT_CODES). "Tir lliure
# convertit/fallat" és el text d'algun partit antic de l'historic.
_RE = {
    "t1c": r"^(Cistella de 1|Tir lliure convertit)",
    "t1f": r"^(Intent fallat de 1|Tir lliure fallat)",
    "t2c": r"^Cistella de 2",
    "t2f": r"^Intent fallat de 2",
    "t3c": r"^Cistella de 3",
    "t3f": r"^Intent fallat de 3",
}

# (clau, etiqueta, unitat) — l'ordre és el de la taula de l'app
INDICADORS = [
    ("pts", "Punts totals", ""),
    ("pts2", "Punts de 2", ""),
    ("pts3", "Punts de 3", ""),
    ("ptl", "Punts de tir lliure", ""),
    ("pct2", "% encert tir de 2", "%"),
    ("pct3", "% encert tir de 3", "%"),
    ("pctl", "% encert tir lliure", "%"),
    ("pctc", "% encert tirs de camp", "%"),
    ("sel3", "% d'intents de 3 sobre tirs de camp", "%"),
    ("tli", "Tirs lliures intentats", ""),
    ("tl_tc", "Tirs lliures intentats / tirs de camp intentats", ""),
    ("t2i", "Tirs de 2 intentats", ""),
    ("t3i", "Tirs de 3 intentats", ""),
]


def _pct(c, i):
    return 100.0 * c / i if i else 0.0


def stats_de_comptes(t2c, t2i, t3c, t3i, tlc, tli):
    """Indicadors d'un equip en un partit a partir d'anotats/intentats."""
    tcc, tci = t2c + t3c, t2i + t3i
    return {
        "t2c": t2c, "t2i": t2i, "t3c": t3c, "t3i": t3i, "tlc": tlc, "tli": tli,
        "tcc": tcc, "tci": tci,
        "pts2": 2 * t2c, "pts3": 3 * t3c, "ptl": tlc,
        "pts": 2 * t2c + 3 * t3c + tlc,
        "pct2": _pct(t2c, t2i), "pct3": _pct(t3c, t3i),
        "pctl": _pct(tlc, tli), "pctc": _pct(tcc, tci),
        "sel3": _pct(t3i, tci),
        "tl_tc": tli / tci if tci else 0.0,
    }


def stats_equip(df, id_equip):
    """Indicadors d'un equip en un partit, comptats al play-by-play."""
    accio = df.loc[df["idEquip"].astype(str) == str(id_equip), "accio"].fillna("")
    n = {k: int(accio.str.match(p, case=False).sum()) for k, p in _RE.items()}
    st = stats_de_comptes(n["t2c"], n["t2c"] + n["t2f"], n["t3c"], n["t3c"] + n["t3f"],
                          n["t1c"], n["t1c"] + n["t1f"])
    # Comprovació: els punts dels tirs han de quadrar amb els del play-by-play
    if "punts" in df.columns:
        pts_pbp = int(pd.to_numeric(df.loc[df["idEquip"].astype(str) == str(id_equip), "punts"],
                                    errors="coerce").fillna(0).sum())
        st["avis_punts"] = None if pts_pbp == st["pts"] else (
            f"els tirs donen {st['pts']} punts però el play-by-play en suma {pts_pbp}")
    return st


def equips_del_partit(df):
    """IDs d'equip d'un partit (sense les files sense equip)."""
    ids = df["idEquip"].astype(str)
    return [i for i in ids.unique() if i not in ("", "0", "nan", "None")]


def identifica_rival(df_vs, dfs_altres):
    """El rival és l'equip del partit contra nosaltres que també surt a tots els
    altres partits. Retorna (id_rival, id_nostre)."""
    eq_vs = equips_del_partit(df_vs)
    if len(eq_vs) != 2:
        raise ValueError("El partit contra nosaltres no té dos equips.")
    comuns = set(eq_vs)
    for d in dfs_altres:
        comuns &= set(equips_del_partit(d))
    if len(comuns) != 1:
        raise ValueError("No es troba cap equip que surti a tots els partits del rival. "
                         "Comprova que els 'altres partits' siguin del mateix rival.")
    id_rival = comuns.pop()
    id_nostre = [e for e in eq_vs if e != id_rival][0]
    return id_rival, id_nostre


def _mitjana(llista):
    claus = [k for k, _, _ in INDICADORS]
    return {k: sum(s[k] for s in llista) / len(llista) for k in claus}


def analitza_rival(nom, df_vs, dfs_altres):
    """{nom, id_rival, id_nostre, vs, resta, partits_resta, avisos}.
    'vs' és el partit contra nosaltres; 'resta' la mitjana dels altres."""
    if not dfs_altres:
        raise ValueError(f"{nom}: cal almenys un altre partit del rival.")
    id_rival, id_nostre = identifica_rival(df_vs, dfs_altres)
    vs = stats_equip(df_vs, id_rival)
    altres = [stats_equip(d, id_rival) for d in dfs_altres]
    avisos = [f"{nom} (contra nosaltres): {vs['avis_punts']}"] if vs.get("avis_punts") else []
    avisos += [f"{nom} (partit {i + 1} de la resta): {s['avis_punts']}"
               for i, s in enumerate(altres) if s.get("avis_punts")]
    return {"nom": nom, "id_rival": id_rival, "id_nostre": id_nostre,
            "vs": vs, "resta": _mitjana(altres), "partits_resta": len(altres),
            "avisos": avisos}


def resum(rivals):
    """Mitjana entre rivals (mateix pes per a tots) de 'vs' i 'resta'."""
    return {"vs": _mitjana([r["vs"] for r in rivals]),
            "resta": _mitjana([r["resta"] for r in rivals])}


def taula(rivals):
    """DataFrame per mostrar a l'app: una fila per indicador."""
    mit = resum(rivals)
    files = []
    for k, etiqueta, _ in INDICADORS:
        f = {"Indicador": etiqueta}
        for r in rivals:
            f[f"{r['nom']} vs"] = r["vs"][k]
            f[f"{r['nom']} resta"] = r["resta"][k]
        f["Mitjana vs"] = mit["vs"][k]
        f["Mitjana resta"] = mit["resta"][k]
        f["Diferència"] = mit["vs"][k] - mit["resta"][k]
        files.append(f)
    return pd.DataFrame(files)


# ══════════════════════════════════════════════════
# TEXTOS (títols i conclusions amb xifres)
# ══════════════════════════════════════════════════

def _n(v, dec=1):
    """Número a la catalana: 30,3"""
    return f"{v:.{dec}f}".replace(".", ",")


def _p(v):
    return f"{v:.0f}%"


def _quants(rivals, clau, signe):
    """Quants rivals van en la mateixa direcció (signe +1 = més contra nosaltres)."""
    return sum(1 for r in rivals if (r["vs"][clau] - r["resta"][clau]) * signe > 0)


def _llista_noms(noms):
    return noms[0] if len(noms) == 1 else ", ".join(noms[:-1]) + " i " + noms[-1]


def textos(rivals, nostre):
    """Tots els textos de l'informe en un dict, perquè l'app els pugui ensenyar
    i el PowerPoint els faci servir igual."""
    m = resum(rivals)
    v, r = m["vs"], m["resta"]
    n = len(rivals)
    d = {k: v[k] - r[k] for k in v}

    def mes_menys(x, llindar=0.0):
        return "més" if x > llindar else ("menys" if x < -llindar else "gairebé igual")

    def millor_pitjor(x):
        return "millor" if x > 0 else "pitjor"

    def cons(clau):
        s = 1 if d[clau] > 0 else -1
        k = _quants(rivals, clau, s)
        return f"{k} de {n} rivals"

    t = {}
    # Punts
    if abs(d["pts"]) < 2:
        t["punts"] = (f"els rivals anoten gairebé igual contra {nostre} "
                      f"({_n(v['pts'])} vs {_n(r['pts'])}, {'+' if d['pts'] >= 0 else ''}{_n(d['pts'])}).")
        t["tit_punts"] = f"Contra {nostre} els rivals anoten gairebé igual ({_n(v['pts'])} vs {_n(r['pts'])})"
    else:
        t["punts"] = (f"els rivals anoten {_n(abs(d['pts']))} punts {mes_menys(d['pts'])} contra {nostre} "
                      f"({_n(v['pts'])} vs {_n(r['pts'])}; {cons('pts')}).")
        t["tit_punts"] = (f"Contra {nostre} els rivals anoten {_n(abs(d['pts']))} punts "
                          f"{mes_menys(d['pts'])} per partit ({_n(v['pts'])} vs {_n(r['pts'])})")

    # Tir de 3
    t["tir3"] = (f"contra {nostre} els rivals tiren {millor_pitjor(d['pct3'])} de 3 "
                 f"({_p(v['pct3'])} vs {_p(r['pct3'])}) i en proporció en tiren "
                 f"{mes_menys(d['sel3'], 0.5)} ({_p(v['sel3'])} vs {_p(r['sel3'])} dels tirs de camp).")
    # Tir lliure
    t["tl"] = (f"els rivals llancen {_n(abs(d['tli']))} tirs lliures {mes_menys(d['tli'])} contra {nostre} "
               f"({_n(v['tli'])} vs {_n(r['tli'])} intents; {cons('tli')}), amb un encert del "
               f"{_p(v['pctl'])} vs {_p(r['pctl'])}.")
    # Joc interior
    t["interior"] = (f"{_n(abs(d['t2i']))} tirs de 2 intentats {mes_menys(d['t2i'])} "
                     f"({_n(v['t2i'])} vs {_n(r['t2i'])}) i {_n(abs(d['pts2']))} punts de 2 "
                     f"{mes_menys(d['pts2'])} ({_n(v['pts2'])} vs {_n(r['pts2'])}), "
                     f"amb un encert del {_p(v['pct2'])} vs {_p(r['pct2'])}.")

    # Títols de les diapositives de gràfic
    tipus = {"pct2": "el tir de 2", "pct3": "el tir de 3", "pctl": "el tir lliure",
             "pctc": "els tirs de camp"}
    gran = max(tipus, key=lambda k: abs(d[k]))
    t["tit_encert"] = (f"La diferència més gran és en {tipus[gran]}: "
                       f"{'+' if d[gran] >= 0 else '−'}{_n(abs(d[gran]))} punts percentuals "
                       f"({_p(v[gran])} vs {_p(r[gran])})")

    def quota(s, k):
        return 100.0 * s[k] / s["pts"] if s["pts"] else 0.0
    q = {k: quota(v, k) - quota(r, k) for k in ("pts2", "pts3", "ptl")}
    nom_q = {"pts2": "tirs de 2", "pts3": "triples", "ptl": "tirs lliures"}
    kq = max(q, key=lambda k: abs(q[k]))
    t["tit_distrib"] = (f"Contra {nostre} pesen {mes_menys(q[kq])} els punts de {nom_q[kq]}: "
                        f"el {_p(quota(v, kq))} dels punts (vs {_p(quota(r, kq))})")
    t["tit_seleccio"] = (f"Contra {nostre} els rivals tiren {mes_menys(d['sel3'], 0.5)} de 3: el "
                         f"{_p(v['sel3'])} dels tirs de camp (vs {_p(r['sel3'])})")
    t["tit_tl"] = (f"Contra {nostre} els rivals llancen {_n(abs(d['tli']))} tirs lliures "
                   f"{mes_menys(d['tli'])} per partit ({_n(v['tli'])} vs {_n(r['tli'])})")
    t["tit_t2"] = (f"Contra {nostre} els rivals intenten {_n(abs(d['t2i']))} tirs de 2 "
                   f"{mes_menys(d['t2i'])} per partit ({_n(v['t2i'])} vs {_n(r['t2i'])})")
    t["tit_t3"] = (f"Contra {nostre} els rivals intenten {_n(abs(d['t3i']))} triples "
                   f"{mes_menys(d['t3i'])} per partit ({_n(v['t3i'])} vs {_n(r['t3i'])})")

    # Conclusions
    pr = sorted({x["partits_resta"] for x in rivals})
    pr_txt = f"{pr[0]}" if len(pr) == 1 else f"{pr[0]}–{pr[-1]}"
    t["c_faltes"] = (f"contra {nostre} els rivals llancen {mes_menys(d['tli'])} tirs lliures "
                     f"({_n(v['tli'])} vs {_n(r['tli'])}; {_n(v['tl_tc'], 2)} vs {_n(r['tl_tc'], 2)} "
                     f"per tir de camp) i n'encerten el {_p(v['pctl'])} vs {_p(r['pctl'])}.")
    t["c_exterior"] = (f"tiren {millor_pitjor(d['pct3'])} de 3 ({_p(v['pct3'])} vs {_p(r['pct3'])}, "
                       f"{cons('pct3')}) i el triple pesa el {_p(v['sel3'])} dels seus tirs de camp "
                       f"(vs {_p(r['sel3'])}).")
    t["c_pintura"] = (f"intenten {_n(v['t2i'])} tirs de 2 per partit (vs {_n(r['t2i'])}) "
                      f"i n'encerten el {_p(v['pct2'])} (vs {_p(r['pct2'])}).")
    t["c_mostra"] = (f"només hi ha {n} rival{'s' if n != 1 else ''} i {pr_txt} partit"
                     f"{'s' if pr[-1] != 1 else ''} de comparació per equip; cal prendre-ho com a "
                     f"tendència i confirmar-ho amb més partits.")
    return t


# ══════════════════════════════════════════════════
# POWERPOINT
# ══════════════════════════════════════════════════

VERMELL = RGBColor(0xC0, 0x00, 0x00)
GRIS = RGBColor(0xA6, 0xA6, 0xA6)
TEXT = RGBColor(0x1A, 0x1A, 0x1A)
GRIS_TXT = RGBColor(0x59, 0x59, 0x59)
BLAU = RGBColor(0x1F, 0x3B, 0x57)
# Un color per rival al gràfic d'encert (la mitjana va en vermell)
COLORS_RIVALS = [RGBColor(0x1F, 0x3B, 0x57), RGBColor(0x5B, 0x8D, 0xB8),
                 RGBColor(0x7F, 0x7F, 0x7F), RGBColor(0xB8, 0xB8, 0xB8),
                 RGBColor(0x2E, 0x7D, 0x6B), RGBColor(0x8E, 0x6C, 0x3A)]
FONT = "Calibri"

W, H = Inches(13.333), Inches(7.5)


def _caixa(slide, x, y, w, h, runs, size, color=TEXT, bold=False, align=PP_ALIGN.LEFT,
           anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    for txt, o in (runs if isinstance(runs, list) else [(runs, {})]):
        r = p.add_run()
        r.text = txt
        f = r.font
        f.name, f.size = FONT, Pt(o.get("size", size))
        f.bold = o.get("bold", bold)
        f.color.rgb = o.get("color", color)
    return tb


def _rect(slide, x, y, w, h, color):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = color
    sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def _vinyetes(slide, x, y, w, h, punts):
    """Vinyetes vermelles a 20pt; el text abans dels dos punts va en negreta."""
    tf = slide.shapes.add_textbox(x, y, w, h).text_frame
    tf.word_wrap = True
    for i, (cap, cos) in enumerate(punts):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", str(Emu(Inches(0.35))))
        pPr.set("indent", str(-Emu(Inches(0.35))))
        clr = pPr.makeelement(qn("a:buClr"), {})
        srgb = clr.makeelement(qn("a:srgbClr"), {"val": "C00000"})
        clr.append(srgb)
        pPr.append(clr)
        pPr.append(pPr.makeelement(qn("a:buFont"), {"typeface": "Arial"}))
        pPr.append(pPr.makeelement(qn("a:buChar"), {"char": "•"}))
        p.space_after = Pt(14)
        for txt, bold in ((cap + ": ", True), (cos[0].upper() + cos[1:], False)):
            r = p.add_run()
            r.text = txt
            r.font.name, r.font.size, r.font.bold = FONT, Pt(20), bold
            r.font.color.rgb = TEXT


def _capcalera(slide, etiqueta, titol):
    """Barreta vermella + etiqueta del tema (14pt) + títol amb la conclusió (26pt)."""
    _rect(slide, Inches(0.6), Inches(0.55), Inches(0.08), Inches(0.3), VERMELL)
    _caixa(slide, Inches(0.8), Inches(0.55), Inches(11.5), Inches(0.3), etiqueta, 14,
           VERMELL, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    _caixa(slide, Inches(0.6), Inches(0.95), Inches(12.1), Inches(0.95), titol, 26,
           TEXT, bold=True)


def _estil_grafic(chart, llegenda=True, etiquetes=True, fmt='0.0', apilat=False):
    chart.font.name = FONT
    chart.font.size = Pt(12)
    chart.has_legend = llegenda
    if llegenda:
        chart.legend.position = XL_LEGEND_POSITION.TOP
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(12)
    va = chart.value_axis
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = RGBColor(0xE5, 0xE5, 0xE5)
    va.format.line.fill.background()
    va.tick_labels.font.size = Pt(11)
    va.tick_labels.font.color.rgb = GRIS_TXT
    va.tick_labels.number_format, va.tick_labels.number_format_is_linked = '0', False
    ca = chart.category_axis
    ca.tick_labels.font.size = Pt(12)
    ca.format.line.color.rgb = RGBColor(0xBF, 0xBF, 0xBF)
    # Amb valors negatius, les etiquetes van a baix i no enmig de les barres
    ca.tick_label_position = XL_TICK_LABEL_POSITION.LOW
    plot = chart.plots[0]
    plot.gap_width = 60 if not apilat else 50
    if not apilat:
        plot.overlap = -10
    if etiquetes:
        plot.has_data_labels = True
        dl = plot.data_labels
        dl.number_format, dl.number_format_is_linked = fmt, False
        dl.font.size = Pt(11)
        dl.position = XL_LABEL_POSITION.CENTER if apilat else XL_LABEL_POSITION.OUTSIDE_END


def _afegeix_grafic(slide, tipus, dades):
    # Zona del gràfic: aprox. de 140 a 512 pt d'alçada, centrat
    return slide.shapes.add_chart(tipus, Inches(0.9), Pt(140), Inches(11.5), Pt(372), dades).chart


def _colors_series(chart, colors):
    for s, c in zip(chart.series, colors):
        s.format.fill.solid()
        s.format.fill.fore_color.rgb = c
        s.invert_if_negative = False


def _grafic_vs_resta(slide, rivals, clau, nostre, fmt='0.0'):
    """Columnes agrupades: contra nosaltres (vermell) i resta (gris), per rival + mitjana."""
    m = resum(rivals)
    cd = CategoryChartData(number_format=fmt)
    cd.categories = [r["nom"] for r in rivals] + ["Mitjana"]
    cd.add_series(f"Contra {nostre}", [r["vs"][clau] for r in rivals] + [m["vs"][clau]])
    cd.add_series("Resta de partits", [r["resta"][clau] for r in rivals] + [m["resta"][clau]])
    ch = _afegeix_grafic(slide, XL_CHART_TYPE.COLUMN_CLUSTERED, cd)
    _estil_grafic(ch, fmt=fmt)
    _colors_series(ch, [VERMELL, GRIS])
    return ch


def genera_pptx(rivals, nostre):
    """Retorna els bytes del .pptx de 10 diapositives."""
    t = textos(rivals, nostre)
    m = resum(rivals)
    noms = [r["nom"] for r in rivals]

    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    blank = prs.slide_layouts[6]

    def nova():
        s = prs.slides.add_slide(blank)
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        return s

    # 1 · Portada
    s = nova()
    _rect(s, Inches(0.9), Inches(2.5), Inches(0.12), Inches(2.3), VERMELL)
    _caixa(s, Inches(1.3), Inches(2.4), Inches(11), Inches(1.5),
           f"Què els passa als rivals quan juguen contra {nostre}?", 40, TEXT, bold=True,
           anchor=MSO_ANCHOR.BOTTOM)
    _caixa(s, Inches(1.3), Inches(4.05), Inches(11), Inches(0.9),
           f"Comparació del partit contra {nostre} amb la mitjana de la resta de partits "
           f"({_llista_noms(noms)})", 20, GRIS_TXT)

    # 2 · Resum
    s = nova()
    _capcalera(s, "Resum", "Resum dels resultats")
    _vinyetes(s, Inches(0.9), Inches(2.1), Inches(11.6), Inches(4.8), [
        ("Punts", t["punts"]), ("Tir de 3", t["tir3"]),
        ("Tir lliure", t["tl"]), ("Joc interior", t["interior"])])

    # 3 · Punts anotats
    s = nova()
    _capcalera(s, "Punts anotats", t["tit_punts"])
    _grafic_vs_resta(s, rivals, "pts", nostre)

    # 4 · Variació del % d'encert (punts percentuals)
    s = nova()
    _capcalera(s, "% d'encert", t["tit_encert"])
    tipus = [("pct2", "Tir de 2"), ("pct3", "Tir de 3"), ("pctl", "Tir lliure"), ("pctc", "Tirs de camp")]
    cd = CategoryChartData(number_format='+0.0;-0.0;0.0')
    cd.categories = [e for _, e in tipus]
    for r in rivals:
        cd.add_series(r["nom"], [r["vs"][k] - r["resta"][k] for k, _ in tipus])
    cd.add_series("Mitjana", [m["vs"][k] - m["resta"][k] for k, _ in tipus])
    ch = _afegeix_grafic(s, XL_CHART_TYPE.COLUMN_CLUSTERED, cd)
    _estil_grafic(ch, fmt='+0.0;-0.0;0.0')
    _colors_series(ch, COLORS_RIVALS[:len(rivals)] + [VERMELL])

    # 5 · D'on venen els punts
    s = nova()
    _capcalera(s, "Distribució dels punts", t["tit_distrib"])
    cd = CategoryChartData(number_format='0.0')
    cats, p2, p3, pl = [], [], [], []
    for r in rivals + [{"nom": "Mitjana", **m}]:
        for clau, eti in (("vs", f"vs {nostre}"), ("resta", "resta")):
            cats.append(f"{r['nom']}\n{eti}")
            p2.append(r[clau]["pts2"]); p3.append(r[clau]["pts3"]); pl.append(r[clau]["ptl"])
    cd.categories = cats
    cd.add_series("Punts de 2", p2)
    cd.add_series("Punts de 3", p3)
    cd.add_series("Punts de tir lliure", pl)
    ch = _afegeix_grafic(s, XL_CHART_TYPE.COLUMN_STACKED, cd)
    _estil_grafic(ch, apilat=True, fmt='0')
    ch.plots[0].data_labels.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    _colors_series(ch, [BLAU, VERMELL, GRIS])

    # 6-9 · Selecció de tir, tirs lliures, tirs de 2, tirs de 3
    for etiqueta, clau_tit, clau, fmt in (
            ("Selecció de tir · % d'intents de 3 sobre tirs de camp", "tit_seleccio", "sel3", '0'),
            ("Tirs lliures intentats", "tit_tl", "tli", '0.0'),
            ("Tirs de 2 intentats", "tit_t2", "t2i", '0.0'),
            ("Tirs de 3 intentats", "tit_t3", "t3i", '0.0')):
        s = nova()
        _capcalera(s, etiqueta, t[clau_tit])
        _grafic_vs_resta(s, rivals, clau, nostre, fmt=fmt)

    # 10 · Conclusions
    s = nova()
    _capcalera(s, "Conclusions", f"Què ens diuen els rivals de {nostre}")
    _vinyetes(s, Inches(0.9), Inches(2.1), Inches(11.6), Inches(4.8), [
        ("Faltes", t["c_faltes"]), ("Defensa exterior", t["c_exterior"]),
        ("Atac a la pintura", t["c_pintura"]), ("Precaució", t["c_mostra"])])

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def nom_fitxer(nostre):
    net = re.sub(r"[^\w]+", "_", nostre, flags=re.UNICODE).strip("_") or "equip"
    return f"Informe_rivals_{net}.pptx"
