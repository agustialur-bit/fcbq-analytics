"""
Comprova els càlculs de l'informe de rivals (informe_rivals.py).

No toca la xarxa. La referència són els tirs que l'Agustí havia passat a mà a
l'Excel "Anàlisis rivals v2.xlsx" (rivals de Manresa): l'informe ha de donar les
mateixes xifres que el seu full "Anàlisi Manresa" (65,0 vs 61,4 punts, 30,3 vs
19,9 tirs lliures...).

    python -m pytest tests/test_informe_rivals.py -q
"""
import io
import sys
from pathlib import Path

import pandas as pd
import pytest
from pptx import Presentation

ARREL = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ARREL))

import informe_rivals as ir  # noqa: E402

FIXTURES = Path(__file__).parent / "fixtures"

# (t2c, t2i, t3c, t3i, tlc, tli) — primer el partit contra Manresa, després la resta
MANRESA = {
    "Santfeliuenc": [(20, 44, 4, 21, 24, 38), (14, 39, 8, 31, 13, 16),
                     (13, 34, 3, 21, 11, 19), (22, 51, 6, 23, 14, 21)],
    "Sant Gervasi": [(17, 40, 3, 11, 17, 32), (15, 47, 6, 18, 18, 25),
                     (22, 47, 6, 21, 11, 15)],
    "Balaguer": [(14, 52, 6, 16, 13, 21), (14, 42, 2, 13, 11, 19),
                 (15, 31, 5, 14, 15, 23)],
}


def _df_partit(id_a, id_b, comptes_a, comptes_b=(10, 30, 2, 10, 5, 10)):
    """Play-by-play mínim amb els tirs d'un partit entre dos equips."""
    files = []
    for eq, (t2c, t2i, t3c, t3i, tlc, tli) in ((id_a, comptes_a), (id_b, comptes_b)):
        for accio, n, pts in (("Cistella de 2", t2c, 2), ("Intent fallat de 2", t2i - t2c, 0),
                              ("Cistella de 3", t3c, 3), ("Intent fallat de 3", t3i - t3c, 0),
                              ("Cistella de 1", tlc, 1), ("Intent fallat de 1", tli - tlc, 0)):
            files += [{"idEquip": eq, "accio": accio, "punts": pts}] * n
        files.append({"idEquip": eq, "accio": "Personal 2 tirs lliures, 1a falta", "punts": 0})
    files.append({"idEquip": "0", "accio": "Final de període", "punts": 0})
    return pd.DataFrame(files)


@pytest.fixture(scope="module")
def rivals_manresa():
    rivals = []
    for i, (nom, partits) in enumerate(MANRESA.items()):
        id_riv = f"riv{i}"
        df_vs = _df_partit(id_riv, "manresa", partits[0])
        dfs = [_df_partit(id_riv, f"altre{i}{j}", c) for j, c in enumerate(partits[1:])]
        rivals.append(ir.analitza_rival(nom, df_vs, dfs))
    return rivals


def test_tirs_del_fixture_quadren_amb_els_punts():
    df = pd.read_csv(FIXTURES / "partit_2477497.csv", encoding="utf-8",
                     dtype={"idEquip": str, "dorsal": str})
    totals = {"t2c": 0, "t2i": 0, "t3c": 0, "t3i": 0, "tlc": 0, "tli": 0}
    for eq in ir.equips_del_partit(df):
        s = ir.stats_equip(df, eq)
        assert s["avis_punts"] is None
        for k in totals:
            totals[k] += s[k]
    assert totals == {"t2c": 34, "t2i": 67, "t3c": 11, "t3i": 46, "tlc": 22, "tli": 36}


def test_detecta_rival_i_nosaltres(rivals_manresa):
    assert [r["id_rival"] for r in rivals_manresa] == ["riv0", "riv1", "riv2"]
    assert {r["id_nostre"] for r in rivals_manresa} == {"manresa"}
    assert [r["partits_resta"] for r in rivals_manresa] == [3, 2, 2]


def test_rival_que_no_surt_a_tots_els_partits():
    df_vs = _df_partit("a", "nosaltres", (10, 20, 1, 5, 2, 4))
    with pytest.raises(ValueError):
        ir.identifica_rival(df_vs, [_df_partit("b", "c", (10, 20, 1, 5, 2, 4))])


def test_xifres_iguals_que_lexcel(rivals_manresa):
    m = ir.resum(rivals_manresa)
    esperat = {"pts": (65.0, 61.4), "tli": (30.3, 19.9), "pctl": (59.4, 67.6),
               "sel3": (25.8, 31.5), "pct3": (27.9, 26.2), "t2i": (45.3, 41.6)}
    for k, (vs, resta) in esperat.items():
        assert m["vs"][k] == pytest.approx(vs, abs=0.05), k
        assert m["resta"][k] == pytest.approx(resta, abs=0.05), k


def test_textos_amb_xifres(rivals_manresa):
    t = ir.textos(rivals_manresa, "Manresa")
    assert t["tit_tl"] == ("Contra Manresa els rivals llancen 10,4 tirs lliures més per "
                           "partit (30,3 vs 19,9)")
    assert "65,0 vs 61,4" in t["punts"]
    assert "3 rivals" in t["c_mostra"]


def test_pptx_de_10_diapositives(rivals_manresa):
    prs = Presentation(io.BytesIO(ir.genera_pptx(rivals_manresa, "Manresa")))
    assert len(prs.slides) == 10
    grafics = [sh for s in prs.slides for sh in s.shapes if sh.has_chart]
    assert len(grafics) == 7
    assert ir.nom_fitxer("CB Manresa") == "Informe_rivals_CB_Manresa.pptx"
